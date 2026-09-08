"""Verifica archivos XLSX reales contra una API local con datos sintéticos."""
import datetime
import http.client
import http.cookiejar
import io
import json
import os
from pathlib import Path
import re
import subprocess
import tempfile
import urllib.error
import urllib.parse
import urllib.request
import uuid
import zipfile
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

BASE = os.environ.get('EXP_IMCYC_TEST_URL', 'http://127.0.0.1:4187/exp-imcyc/')
assert urllib.parse.urlparse(BASE).hostname in ('127.0.0.1', 'localhost'), 'Usar solo una instancia local de pruebas.'
ORIGIN = 'http://' + urllib.parse.urlparse(BASE).netloc
MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
checks = 0


class LocalConnection(http.client.HTTPConnection):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, source_address=('127.0.0.3', 0), **kwargs)


class LocalHandler(urllib.request.HTTPHandler):
    def http_open(self, request):
        return self.do_open(LocalConnection, request)


class Client:
    def __init__(self):
        self.opener = urllib.request.build_opener(LocalHandler(), urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
        self.csrf = ''

    def call(self, action, body=None, query=None, expected=200):
        global checks
        headers = {'Origin': ORIGIN}
        payload = None
        if body is not None:
            payload = json.dumps(body).encode()
            headers.update({'Content-Type': 'application/json', 'X-CSRF-Token': self.csrf})
        request = urllib.request.Request(BASE + 'api/index.php?' + urllib.parse.urlencode({'action': action, **(query or {})}), data=payload, headers=headers)
        try:
            response = self.opener.open(request)
        except urllib.error.HTTPError as error:
            response = error
        content = response.read()
        assert response.status == expected, (action, response.status, content[:500])
        assert 'no-store' in response.headers.get('Cache-Control', '')
        checks += 1
        if response.headers.get_content_type() == MIME:
            assert re.fullmatch(r'attachment; filename="registros-imcyc-\d{8}-\d{6}\.xlsx"', response.headers['Content-Disposition'])
            assert int(response.headers['Content-Length']) == len(content)
            assert response.headers['X-Content-Type-Options'] == 'nosniff'
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                assert archive.testzip() is None
                assert not any('externalLink' in name or 'vbaProject' in name for name in archive.namelist())
            workbook = load_workbook(io.BytesIO(content), data_only=False)
            assert workbook.sheetnames == ['Registros']
            sheet = workbook.active
            assert sheet.max_row - 1 == int(response.headers['X-Export-Count'])
            assert sheet.max_column == 15 and sheet.freeze_panes == 'B2'
            assert sheet.auto_filter.ref == f'A1:O{sheet.max_row}'
            assert all(cell.data_type != 'f' for row in sheet for cell in row), 'Los participantes no pueden introducir fórmulas.'
            return sheet
        result = json.loads(content)
        self.csrf = result.get('csrf', self.csrf)
        return result


anonymous = Client()
anonymous.call('admin.registrations.export', expected=401)
anonymous.call('admin.registrations.export', body={}, expected=405)
admin = Client()
admin.call('session')
admin.call('login', {'username': 'qa', 'password': os.environ.get('EXP_IMCYC_TEST_PASSWORD', 'Clave-QA-local-2026')})
for query in [{'status': 'invalid'}, {'service': 'invalid'}, {'search[]': 'invalid'}, {'consent': 'invalid'}, {'consent[]': 'yes'}]:
    admin.call('admin.registrations.export', query=query, expected=422)

prefix = 'xlsx-' + uuid.uuid4().hex[:10]
participant = Client()
participant.call('session')
registrations = []
for index in range(27):
    participant.call('registration.reset', {})
    participant.call('registration.save', {
        'name': '=HYPERLINK("https://example.com", "María & José")' if index == 0 else f'Participante Excel {index:02d}',
        'company': f'{prefix} · Construcción <IMCYC> & México',
        'email': f'{prefix}-{index}@example.com',
    })
    if index < 3:
        participant.call('survey.submit', {
            'service': 'Certificación', 'application': '+SUM(1,2)\n' + 'Aplicación técnica & calidad <prueba>. ' * 100,
            'clarity': 'Buena', 'serviceRating': 0,
            **({'opinionConsent': index == 0} if index < 2 else {}),
        })
    registrations.append(participant.call('registration.current')['registration'])
participant.call('admin.registrations.export', expected=401)

sheet = admin.call('admin.registrations.export', query={'search': prefix, 'page': '2'})
assert sheet.max_row == 28, 'Exportar todas las páginas, no solo las 25 filas visibles.'
assert [row[0].value for row in list(sheet.rows)[1:]] == [str(row['id']) for row in reversed(registrations)]
for cells, registration in zip(list(sheet.rows)[1:], reversed(registrations)):
    values = [cell.value for cell in cells]
    assert values[:4] == [str(registration['id']), registration['full_name'], registration['email'], registration['company']]
    assert values[5:10] == [registration[key] for key in ('service', 'application', 'clarity', 'service_rating', 'unique_code')]
    for index, key in [(10, 'created_at'), (11, 'completed_at'), (13, 'opinion_consent_at')]:
        expected = registration[key]
        if expected:
            expected = datetime.datetime.fromisoformat(expected).astimezone(ZoneInfo('America/Mexico_City')).replace(tzinfo=None)
        assert values[index] == expected
    expected_consent = {True: 'Autorizó', False: 'No autorizó', None: 'Sin autorización registrada'}[registration['opinion_consent']]
    assert values[12] == expected_consent and values[14] == registration['opinion_consent_text']
assert sheet['B28'].data_type == 's' and sheet['I28'].data_type == 'n' and sheet['I28'].value == 0
assert sheet['I2'].value is None and sheet['L2'].value is None
assert sheet['A1'].font.bold and sheet['A1'].fill.patternType == 'solid'

pending = admin.call('admin.registrations.export', query={'search': prefix, 'status': 'pending'})
assert pending.max_row == 25 and all(row[4].value == 'Solo datos de contacto' for row in list(pending.rows)[1:])
completed = admin.call('admin.registrations.export', query={'search': prefix, 'status': 'completed', 'service': 'Certificación'})
assert completed.max_row == 4 and completed['E2'].value == 'Encuesta completa' and completed['I2'].value == 0
for value, label, count in [('yes', 'Autorizó', 1), ('no', 'No autorizó', 1), ('unrecorded', 'Sin autorización registrada', 25)]:
    filtered = admin.call('admin.registrations.export', query={'search': prefix, 'consent': value})
    listing = admin.call('admin.registrations', query={'search': prefix, 'consent': value})
    assert listing['total'] == count and filtered.max_row == count + 1
    assert all(row[12].value == label for row in list(filtered.rows)[1:])
assert completed['M2'].value == 'Sin autorización registrada' and completed['N2'].value is None and completed['O2'].value is None
folio = admin.call('admin.registrations.export', query={'search': registrations[0]['unique_code']})
assert folio.max_row == 2 and folio['J2'].value == registrations[0]['unique_code']
for search in [prefix + '-no-existe', "' OR 1=1 --", '%']:
    empty = admin.call('admin.registrations.export', query={'search': search})
    assert empty.max_row == 1 and empty['A1'].value == 'ID de registro'
listing = admin.call('admin.registrations')
all_rows = admin.call('admin.registrations.export')
assert all_rows.max_row - 1 == listing['total']
admin.call('logout', {})
admin.call('admin.registrations.export', expected=401)

# Fechas que cambian de día en CDMX e identificadores que Excel redondearía como números.
with tempfile.TemporaryDirectory(prefix='imcyc-xlsx-fixture-') as directory:
    path = Path(directory) / 'fixture.xlsx'
    subprocess.run(['php', str(Path(__file__).with_name('export-fixture.php')), str(path)], check=True)
    fixture = load_workbook(path).active
    assert fixture['A2'].value == '9007199254740993' and fixture['A2'].data_type == 's'
    assert fixture['K2'].value == datetime.datetime(2026, 9, 7, 19, 30)
    assert fixture['L2'].value == datetime.datetime(2026, 9, 7, 20, 15)
    assert fixture['G2'].value == '=SUM(1,2)\nJosé & María <concreto>' and fixture['G2'].data_type == 's'
    assert fixture['M2'].value == 'Autorizó' and fixture['N2'].value == datetime.datetime(2026, 9, 7, 20, 15)

print(f'{checks} solicitudes verificadas: XLSX válido, acceso, filtros, todas las páginas, Unicode, cero, fechas CDMX, texto sin fórmulas y tres estados de autorización.')
